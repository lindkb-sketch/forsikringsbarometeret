#!/usr/bin/env python3
"""Fail-closed Norway extraction from issuer-linked Gjensidige key-figures XLSX.

Source is the issuer's own reports index. Never infer geography from group figures.
Values are stored in the workbook as NOK (not NOK millions) and ratios as fractions.
"""
import hashlib,io,json,re,urllib.parse,urllib.request
from html.parser import HTMLParser
from pathlib import Path
import openpyxl

INDEX='https://www.gjensidige.com/investor-relations/reports-and-presentations'
PDFS={1:'https://mb.cision.com/Public/1122/4341796/af7aaae92042a970.pdf',
      2:'https://mb.cision.com/Public/1122/4373892/b3fb98691b4814f0.pdf'}
SHEETS=('General Insurance Private Norway','General Insurance Commercial Norway')
LABELS={'insurance_revenue':'Insurance revenue','insurance_service_result':'Insurance service result',
        'loss_ratio':'Loss ratio, net of reinsurance','expense_ratio':'Cost ratio',
        'combined_ratio':'Combined ratio'}
class Links(HTMLParser):
    def __init__(self): super().__init__();self.links=[];self.href=None;self.label=''
    def handle_starttag(self,tag,attrs):
        if tag=='a':self.href=dict(attrs).get('href');self.label=''
    def handle_data(self,data):
        if self.href is not None:self.label+=data
    def handle_endtag(self,tag):
        if tag=='a' and self.href:self.links.append((self.label.strip(),self.href));self.href=None

def get(url):
    req=urllib.request.Request(url,headers={'User-Agent':'Forsikringsbarometeret/1.0'})
    with urllib.request.urlopen(req,timeout=35) as r:
        if r.status!=200:raise ValueError('HTTP status '+str(r.status))
        blob=r.read(20_000_001)
    if len(blob)>20_000_000:raise ValueError('Source too large')
    return blob

def locate_xlsx(page,quarter,base=INDEX):
    """Only direct workbook links anchored in the Qn 2026 key-figures row."""
    from html import unescape
    raw=unescape(page.decode('utf-8','replace'))
    # Dynamic issuer pages sometimes have no anchor in static HTML: refuse to guess.
    candidates=[]
    for match in re.finditer('href="([^"]+[.]xlsx)"',raw,re.I):
        url=urllib.parse.urljoin(base,match.group(1))
        if urllib.parse.urlsplit(url).scheme!='https':continue
        context=raw[max(0,match.start()-1500):match.end()+300].lower()
        if re.search(r'key figures|nøkkeltall',context) and re.search(r'2026',context):
            candidates.append(url)
    candidates=list(dict.fromkeys(candidates))
    # A link's proximity alone cannot identify Q1 vs Q2. Require quarter in URL.
    exact=[u for u in candidates if re.search(r'(?i)(?:q|quarter)[_\s%-]*'+str(quarter)+r'(?!\d)',urllib.parse.unquote(u))]
    if len(exact)!=1:raise ValueError(f'Cannot uniquely verify Q{quarter} 2026 key-figures XLSX URL: {exact}')
    return exact[0]

def parse(blob,url,quarter,year=2026):
    if not blob.startswith(b'PK'):raise ValueError('Not an XLSX archive')
    wb=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    sheets={}
    for prefix in SHEETS:
        matches=[s for s in wb.sheetnames if s.startswith(prefix[:29])]
        if len(matches)!=1:raise ValueError('Missing or ambiguous Norway worksheet: '+prefix)
        sheets[prefix]=wb[matches[0]]
    output=[]
    for period,col in [('quarter',2)]+([('ytd',4)] if quarter>1 else []):
        totals={}
        for prefix,ws in sheets.items():
            header=str(ws.cell(3,col).value or '')
            expected=('Q'+str(quarter) if period=='quarter' else 'YTD')+' '+str(year)
            if header.strip()!=expected:raise ValueError('Unexpected column heading: '+header+' != '+expected)
            rows={}
            for row in range(1,ws.max_row+1):
                label=re.sub(r'[¹²³⁴⁵⁶⁷⁸⁹⁰]+','',str(ws.cell(row,1).value or '')).strip()
                if label in LABELS.values():rows.setdefault(label,[]).append(row)
            for metric,label in LABELS.items():
                if len(rows.get(label,[]))!=1:raise ValueError('Ambiguous metric '+label+' in '+prefix)
                cell=ws.cell(rows[label][0],col)
                if type(cell.value) not in (int,float):raise ValueError('Non-numeric source cell '+cell.coordinate)
                value=float(cell.value)*(100 if 'ratio' in metric else 1/1_000_000)
                totals.setdefault(metric,[]).append((value,ws.title,cell.coordinate))
        revenue=sum(x[0] for x in totals['insurance_revenue'])
        service=sum(x[0] for x in totals['insurance_service_result'])
        loss_expense={}
        for metric in ('loss_ratio','expense_ratio','combined_ratio'):
            loss_expense[metric]=sum(totals['insurance_revenue'][i][0]*totals[metric][i][0] for i in range(2))/revenue
        if abs(loss_expense['combined_ratio']-loss_expense['loss_ratio']-loss_expense['expense_ratio'])>0.02:
            raise ValueError('Norway combined ratio mismatch')
        for metric,value in [('insurance_revenue',revenue),('insurance_service_result',service)]+list(loss_expense.items()):
            output.append(dict(company='Gjensidige',year=year,quarter=quarter,scope='norway',source_scope='norway',
                period_type=period,metric=metric,value=round(value,6),unit='percent' if 'ratio' in metric else 'MNOK',
                status='reported' if metric in ('insurance_revenue','insurance_service_result') else 'calculated',
                calculation=('Revenue-weighted Private Norway + Commercial Norway ratio' if 'ratio' in metric else None),
                source_url=url,source_type='official_investor_tables',source_document_url=PDFS[quarter],
                source_sheet=' + '.join(x[1] for x in totals[metric]),
                source_cell=' + '.join(x[2] for x in totals[metric]),
                source_sha256=hashlib.sha256(blob).hexdigest(),
                source_section='Private Norway + Commercial Norway; NOK -> MNOK; ratios revenue weighted'))
    return output

if __name__=='__main__':
    from audit_financials import DATA,audit
    from build_verified_feed import publish
    page=get(INDEX)
    data=json.loads(DATA.read_text())
    for q in (1,2):
        url=locate_xlsx(page,q)
        rows=parse(get(url),url,q)
        keys=lambda r:tuple(r[k] for k in ('company','year','quarter','scope','period_type','metric'))
        old={keys(r):r for r in data['observations']}
        for r in rows:
            if keys(r) in old and abs(old[keys(r)]['value']-r['value'])>0.001:raise ValueError('Conflicting Norway observation')
            old[keys(r)]=r
        data['observations']=list(old.values())
    if not audit(data)['ok']:raise ValueError('Audit failed')
    DATA.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n')
    publish()
